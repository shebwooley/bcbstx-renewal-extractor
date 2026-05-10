"""
BCBSTX Renewal Extractor, no-AI version.

Usage from Python:
    from bcbstx_extractor import extract_bcbstx_renewal, save_outputs
    data = extract_bcbstx_renewal("renewal.pdf")
    files = save_outputs(data, "renewal", output_dir=".")

Usage from command line:
    python bcbstx_extractor.py renewal.pdf

Dependencies:
    pip install pdfplumber pandas openpyxl
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
import pdfplumber


# ---------- helpers ----------

def parse_money(s: Any) -> float | None:
    if s is None:
        return None
    s = str(s).replace("$", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def normalize_name(name: Any) -> str | None:
    """
    Normalize names so the plan grids match the census:
    - Remove newlines
    - Collapse multiple spaces
    - Normalize spaces around commas
    - Uppercase so case differences do not matter
    """
    if not name:
        return None
    s = str(name).replace("\n", " ")
    s = " ".join(s.split())
    s = re.sub(r"\s*,\s*", ", ", s)
    return s.upper()


def _safe_page_text(pdf: pdfplumber.PDF, page_index: int) -> str:
    if page_index >= len(pdf.pages):
        return ""
    return pdf.pages[page_index].extract_text() or ""


# ---------- group / rating type ----------

def extract_group_info(pdf: pdfplumber.PDF) -> dict[str, Any]:
    text0 = _safe_page_text(pdf, 0)
    text1 = _safe_page_text(pdf, 1)
    group: dict[str, Any] = {}

    m = re.search(r"Account Number:\s*(\d+)\s+Renewal Effective Date:\s*([^\n]+)", text1)
    if m:
        group["account_number"] = m.group(1).strip()
        group["renewal_effective_date"] = m.group(2).strip()

    m2 = re.search(r"\n([A-Z0-9& ,.'/-]+?)\s+Producer:", text1)
    if m2:
        group["group_name"] = m2.group(1).strip()

    m3 = re.search(r"Rating Area:\s*([0-9]+)", text1)
    if m3:
        group["rating_area"] = int(m3.group(1))

    m4 = re.search(r"RateTyp:\s*(\w+)", text0)
    if m4:
        group["rating_type_flag"] = m4.group(1).lower()

    if "RETURN ADDRESS REQUESTED" in text0:
        block = text0.split("RETURN ADDRESS REQUESTED", 1)[1]
        if "Dear Group Administrator" in block:
            block = block.split("Dear Group Administrator", 1)[0]
        lines = [ln.strip() for ln in block.splitlines() if ln.strip() and not set(ln.strip()) == {"*"}]
        if len(lines) >= 3:
            group["mailing_name"] = lines[0]
            group["street"] = lines[1]
            csz = lines[2]
            m = re.match(r"(.+?)\s+([A-Z]{2})\s+(\d{5})", csz)
            if m:
                group["city"] = m.group(1).title()
                group["state"] = m.group(2)
                group["zip"] = m.group(3)
            else:
                group["city_state_zip_raw"] = csz

    return group


def extract_rating_type(pdf: pdfplumber.PDF, group_info: dict[str, Any]) -> str:
    flag = str(group_info.get("rating_type_flag", "")).lower()
    if flag.startswith("composite"):
        return "composite"
    if flag.startswith("age"):
        return "age"

    full_text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    if "Current and Renewal Metallic Medical Plans and Premium - Age Rates" in full_text:
        return "age"
    if "Composite Rates - Medical" in full_text:
        return "composite"
    return "unknown"


# ---------- composite-rated plans ----------

def extract_composite_plans(pdf: pdfplumber.PDF) -> list[dict[str, Any]]:
    plans: list[dict[str, Any]] = []
    for page in pdf.pages:
        for table in page.extract_tables() or []:
            if not table or len(table[0]) < 4:
                continue
            header = table[0]
            if ("Enrolled Count" in (header[1] or "")) and ("Current Plan" in (header[2] or "")):
                plan_id = None
                if len(table) >= 2 and (table[1][0] or "").strip() == "Plan ID":
                    plan_id = (table[1][2] or "").strip()
                if not plan_id:
                    continue

                tier_counts: dict[str, dict[str, Any]] = {}
                current_total = renewal_total = None
                current_eo = renewal_eo = None

                for idx, row in enumerate(table):
                    if idx < 2:
                        continue
                    c0 = (row[0] or "").strip()
                    if c0.startswith("Total Monthly Medical Premium"):
                        current_total = parse_money(row[2])
                        renewal_total = parse_money(row[3])
                        break
                    if len(row) < 4:
                        continue
                    for col_idx in (2, 3):
                        cell = row[col_idx] or ""
                        m = re.search(r"\b(EO|ES|EC|EF):\s*\$?([\d,]+\.\d{2})", cell)
                        if m:
                            tier = m.group(1)
                            rate = parse_money(m.group(2))
                            key = "current_rate" if col_idx == 2 else "renewal_rate"
                            tier_counts.setdefault(tier, {})[key] = rate
                            if row[1]:
                                try:
                                    tier_counts[tier]["count"] = int(row[1])
                                except ValueError:
                                    pass
                            if tier == "EO":
                                if col_idx == 2:
                                    current_eo = rate
                                else:
                                    renewal_eo = rate

                plans.append({
                    "plan_id": plan_id,
                    "rating_basis": "composite",
                    "current_total": current_total,
                    "renewal_total": renewal_total,
                    "current_eo": current_eo,
                    "renewal_eo": renewal_eo,
                    "tier_details": tier_counts,
                })
    return plans


# ---------- age-rated plans ----------

def extract_age_plans(pdf: pdfplumber.PDF) -> list[dict[str, Any]]:
    current_totals: dict[str, float] = {}
    renewal_totals: dict[str, float] = {}

    # Table-based extraction
    for page in pdf.pages:
        for table in page.extract_tables() or []:
            if not table:
                continue
            flat_cells = [(c or "").strip() for row in table for c in row if c]
            table_text = " ".join(flat_cells)
            if "Current Plan ID" not in table_text:
                continue

            for m in re.finditer(r"Current Plan ID[: ]+\s*([A-Z0-9]+)", table_text):
                plan_id = m.group(1)
                if plan_id.startswith("DTX"):
                    continue
                for row in table:
                    row_cells = [(c or "").strip() for c in row]
                    row_text = " ".join(row_cells)
                    if "Total Monthly Medical Premium" not in row_text:
                        continue
                    amounts = []
                    for cell in row_cells:
                        for amt_match in re.findall(r"\$?\d[\d,]*\.\d{2}", cell):
                            val = parse_money(amt_match)
                            if val is not None:
                                amounts.append(val)
                    if amounts:
                        current_totals.setdefault(plan_id, amounts[0])
                        if len(amounts) > 1:
                            renewal_totals.setdefault(plan_id, amounts[1])
                    break

    # Fallback text-based extraction
    for page in pdf.pages:
        text = page.extract_text() or ""
        if "Current Plan ID" not in text or "Total Monthly Medical Premium" not in text:
            continue
        for m in re.finditer(r"Current Plan ID[: ]+\s*([A-Z0-9]+)", text):
            plan_id = m.group(1)
            if plan_id.startswith("DTX"):
                continue
            if plan_id in current_totals and plan_id in renewal_totals:
                continue
            subseq = text[m.start():]
            mtot = re.search(r"Total Monthly Medical Premium[^\n]*", subseq)
            if not mtot:
                continue
            line = mtot.group(0)
            amounts = []
            for amt_match in re.findall(r"\$?\d[\d,]*\.\d{2}", line):
                val = parse_money(amt_match)
                if val is not None:
                    amounts.append(val)
            if amounts:
                current_totals.setdefault(plan_id, amounts[0])
                if len(amounts) > 1:
                    renewal_totals.setdefault(plan_id, amounts[1])

    plans = []
    for pid in sorted(set(current_totals) | set(renewal_totals)):
        plans.append({
            "plan_id": pid,
            "plan_name": None,
            "rating_basis": "age",
            "enrolled_count": None,
            "current_total": current_totals.get(pid),
            "renewal_total": renewal_totals.get(pid),
        })
    return plans


def extract_age_plan_members(pdf: pdfplumber.PDF) -> dict[str, list[dict[str, Any]]]:
    plan_members: dict[str, list[dict[str, Any]]] = {}
    for page in pdf.pages:
        for table in page.extract_tables() or []:
            if not table or len(table) < 3:
                continue
            flat_cells = " ".join((c or "").strip() for row in table for c in row if c)
            if "Current Plan ID" not in flat_cells:
                continue
            m = re.search(r"Current Plan ID[: ]+\s*([A-Z0-9]+)", flat_cells)
            if not m:
                continue
            plan_id = m.group(1)
            if plan_id.startswith("DTX"):
                continue
            members = []
            for row in table[2:]:
                row_text = " ".join((c or "") for c in row)
                if "Total Monthly Medical Premium" in row_text:
                    break
                first = (row[0] or "").strip()
                if not first or not first[0].isdigit():
                    continue
                raw_name = row[1] or ""
                dob = (row[2] or "").strip()
                if not raw_name or not dob:
                    continue
                members.append({"name_norm": normalize_name(raw_name), "dob": dob})
            if members:
                plan_members.setdefault(plan_id, []).extend(members)
    return plan_members


def map_age_plans_to_census(census_rows: list[dict[str, Any]], plan_members: dict[str, list[dict[str, Any]]]) -> dict[int, str]:
    employee_index: dict[tuple[str | None, str], int] = {}
    for row in census_rows:
        relationship = str(row.get("relationship", "")).strip().lower()
        if relationship == "employee":
            key = (normalize_name(row.get("name")), (row.get("dob") or "").strip())
            employee_index[key] = row.get("family_id")

    family_plan: dict[int, str] = {}
    for plan_id, members in plan_members.items():
        for member in members:
            key = (member["name_norm"], member["dob"])
            fam_id = employee_index.get(key)
            if fam_id is not None:
                family_plan.setdefault(fam_id, plan_id)

    for row in census_rows:
        fam_id = row.get("family_id")
        row["plan_id"] = family_plan.get(fam_id)
    return family_plan


# ---------- S662CHC EO composite ----------

def extract_s662_eo(pdf: pdfplumber.PDF) -> float | None:
    for page in pdf.pages:
        for table in page.extract_tables() or []:
            if not table:
                continue
            for idx, row in enumerate(table):
                row_text = " ".join((c or "") for c in row)
                if "Plan ID: S662CHC" in row_text:
                    for r in table[idx + 1:]:
                        if (r[0] or "").strip() == "EO":
                            return parse_money(r[1])
    return None


# ---------- census ----------

def extract_census(pdf: pdfplumber.PDF) -> tuple[dict[str, int], list[dict[str, Any]]]:
    census_counts: dict[str, int] = {}
    census_rows: list[dict[str, Any]] = []

    for page in pdf.pages:
        for tbl in page.extract_tables() or []:
            if not tbl:
                continue
            header = tbl[0]
            if len(header) >= 2 and header[0] == "Tier" and header[1] == "Count":
                if len(tbl) >= 2:
                    row = tbl[1]
                    for i in range(0, len(row) - 1, 2):
                        tier = row[i]
                        cnt = row[i + 1]
                        if tier and cnt:
                            try:
                                census_counts[tier] = int(cnt)
                            except ValueError:
                                pass
                continue

            if len(header) >= 7 and header[0] == "Row" and "Coverage Type" in (header[5] or ""):
                for row in tbl[1:]:
                    if not row or not row[0]:
                        continue
                    fam_str = row[0]
                    try:
                        fam_id = int(str(fam_str).split(".")[0])
                    except ValueError:
                        fam_id = None
                    census_rows.append({
                        "row_id": fam_str,
                        "family_id": fam_id,
                        "name": row[1],
                        "relationship": row[2],
                        "dob": row[3],
                        "age": int(row[4]) if row[4] and str(row[4]).isdigit() else None,
                        "coverage_type": row[5],
                        "state": row[6],
                    })
    return census_counts, census_rows


# ---------- main extraction ----------

def extract_bcbstx_renewal(pdf_path: str | Path) -> dict[str, Any]:
    pdf_path = Path(pdf_path)
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF not found: {pdf_path}")

    with pdfplumber.open(str(pdf_path)) as pdf:
        group = extract_group_info(pdf)
        rating_type = extract_rating_type(pdf, group)

        data: dict[str, Any] = {
            "source_file": pdf_path.name,
            "group_info": group,
            "rating_type": rating_type,
        }

        if rating_type == "composite":
            data["plans"] = extract_composite_plans(pdf)
        elif rating_type == "age":
            data["plans"] = extract_age_plans(pdf)
        else:
            data["plans"] = []

        data["s662chc_eo_composite"] = extract_s662_eo(pdf)
        census_counts, census_rows = extract_census(pdf)

        if rating_type == "age":
            age_plan_members = extract_age_plan_members(pdf)
            family_plan_map = map_age_plans_to_census(census_rows, age_plan_members)
            data["age_plan_members"] = age_plan_members
            data["family_plan_map"] = family_plan_map

        data["census_tier_counts"] = census_counts
        data["census_rows"] = census_rows

    return data


# ---------- output helpers ----------

def build_plan_summary_df(result: dict[str, Any]) -> pd.DataFrame:
    rows = []
    for p in result.get("plans", []):
        cur_tot = p.get("current_total")
        ren_tot = p.get("renewal_total")
        row = {
            "Plan ID": p.get("plan_id"),
            "Rating Basis": p.get("rating_basis"),
            "Current Total": cur_tot,
            "Renewal Total": ren_tot,
            "Total % Change": ((ren_tot - cur_tot) / cur_tot) if cur_tot else None,
        }
        if p.get("rating_basis") == "composite":
            td = p.get("tier_details") or {}
            cur_eo = p.get("current_eo")
            ren_eo = p.get("renewal_eo")
            row.update({
                "Current EO": cur_eo,
                "Renewal EO": ren_eo,
                "EO % Change": ((ren_eo - cur_eo) / cur_eo) if cur_eo else None,
                "Enrolled EO": (td.get("EO") or {}).get("count"),
                "Enrolled ES": (td.get("ES") or {}).get("count"),
                "Enrolled EC": (td.get("EC") or {}).get("count"),
                "Enrolled EF": (td.get("EF") or {}).get("count"),
            })
        rows.append(row)
    return pd.DataFrame(rows)


def build_tier_details_df(result: dict[str, Any]) -> pd.DataFrame:
    rows = []
    for p in result.get("plans", []):
        for tier, d in (p.get("tier_details") or {}).items():
            cur = d.get("current_rate")
            ren = d.get("renewal_rate")
            rows.append({
                "Plan ID": p.get("plan_id"),
                "Tier": tier,
                "Count": d.get("count"),
                "Current Rate": cur,
                "Renewal Rate": ren,
                "% Change": ((ren - cur) / cur) if cur else None,
            })
    return pd.DataFrame(rows)


def build_census_counts_df(result: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame([{"Tier": k, "Count": v} for k, v in (result.get("census_tier_counts") or {}).items()])


def build_age_plan_members_df(result: dict[str, Any]) -> pd.DataFrame:
    rows = []
    for plan_id, members in (result.get("age_plan_members") or {}).items():
        for m in members:
            rows.append({"Plan ID": plan_id, "Member Name Normalized": m.get("name_norm"), "DOB": m.get("dob")})
    return pd.DataFrame(rows)


def build_family_plan_map_df(result: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame([{"Family ID": k, "Plan ID": v} for k, v in (result.get("family_plan_map") or {}).items()])


def build_validation_notes(result: dict[str, Any]) -> pd.DataFrame:
    notes = []
    if not result.get("group_info", {}).get("group_name"):
        notes.append({"Severity": "Warning", "Message": "Group name was not found."})
    if result.get("rating_type") == "unknown":
        notes.append({"Severity": "Warning", "Message": "Rating type could not be identified."})
    if not result.get("plans"):
        notes.append({"Severity": "Warning", "Message": "No medical plans were extracted."})
    if not result.get("census_rows"):
        notes.append({"Severity": "Warning", "Message": "No census rows were extracted."})

    for p in result.get("plans", []):
        if p.get("current_total") is None:
            notes.append({"Severity": "Warning", "Message": f"Current total missing for {p.get('plan_id')}"})
        if p.get("renewal_total") is None:
            notes.append({"Severity": "Warning", "Message": f"Renewal total missing for {p.get('plan_id')}"})

    if result.get("rating_type") == "age":
        census_rows = result.get("census_rows") or []
        family_ids = sorted({r.get("family_id") for r in census_rows if r.get("family_id") is not None})
        missing = []
        for fid in family_ids:
            fam_rows = [r for r in census_rows if r.get("family_id") == fid]
            if fam_rows and not fam_rows[0].get("plan_id"):
                missing.append(fid)
        if missing:
            notes.append({"Severity": "Warning", "Message": "Some age-rated families have no plan_id: " + ", ".join(map(str, missing))})

    if not notes:
        notes.append({"Severity": "OK", "Message": "No validation warnings."})
    return pd.DataFrame(notes)


def autosize_excel(path: str | Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            vals = [str(c.value) for c in col if c.value is not None]
            width = min(max([len(v) for v in vals] + [10]) + 2, 55)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
            header = str(ws.cell(row=1, column=col[0].column).value or "")
            for cell in col[1:]:
                if isinstance(cell.value, (int, float)):
                    if "%" in header:
                        cell.number_format = "0.0%"
                    elif any(x in header for x in ["Total", "Rate", "EO"]):
                        cell.number_format = "$#,##0.00"
    wb.save(path)


def save_outputs(data: dict[str, Any], base_filename: str = "bcbstx_extract", output_dir: str | Path = ".") -> dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_base = re.sub(r"[^A-Za-z0-9_.-]+", "_", Path(base_filename).stem).strip("_") or "bcbstx_extract"

    json_path = output_dir / f"{safe_base}_extract.json"
    xlsx_path = output_dir / f"{safe_base}_broker_review.xlsx"

    with json_path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)

    summary = data.get("group_info", {}).copy()
    summary["rating_type"] = data.get("rating_type")
    summary["s662chc_eo_composite"] = data.get("s662chc_eo_composite")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        pd.DataFrame([summary]).to_excel(writer, sheet_name="Summary", index=False)
        build_plan_summary_df(data).to_excel(writer, sheet_name="Plan Summary", index=False)

        optional_sheets = [
            (build_tier_details_df(data), "Tier Details"),
            (pd.DataFrame(data.get("census_rows") or []), "Census"),
            (build_census_counts_df(data), "Census Counts"),
            (build_age_plan_members_df(data), "Age Plan Members"),
            (build_family_plan_map_df(data), "Family Plan Map"),
        ]
        for df, sheet_name in optional_sheets:
            if not df.empty:
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        build_validation_notes(data).to_excel(writer, sheet_name="Validation Notes", index=False)

    autosize_excel(xlsx_path)
    return {"json_path": json_path, "excel_path": xlsx_path}


def run_extraction(pdf_path: str | Path, output_dir: str | Path = ".") -> dict[str, Any]:
    pdf_path = Path(pdf_path)
    data = extract_bcbstx_renewal(pdf_path)
    output_paths = save_outputs(data, pdf_path.stem, output_dir=output_dir)
    return {"data": data, **output_paths}


# ---------- CLI ----------

def main() -> None:
    parser = argparse.ArgumentParser(description="Extract BCBSTX renewal PDF data into JSON and Excel, no AI.")
    parser.add_argument("pdf_path", help="Path to the BCBSTX renewal PDF")
    parser.add_argument("--output-dir", default=".", help="Folder where output files should be saved")
    args = parser.parse_args()

    result = run_extraction(args.pdf_path, args.output_dir)
    data = result["data"]
    plans = data.get("plans", [])
    validation = build_validation_notes(data)

    print("\n=== BCBSTX Renewal Extractor ===")
    print("Rating basis:", str(data.get("rating_type", "unknown")).upper())
    print("Plans detected:", ", ".join([p.get("plan_id") for p in plans if p.get("plan_id")]) or "NONE")
    print("Census rows:", len(data.get("census_rows") or []))
    print("Census tier counts:", data.get("census_tier_counts"))
    print("\nValidation notes:")
    print(validation.to_string(index=False))
    print("\nFiles created:")
    print("JSON: ", result["json_path"])
    print("Excel:", result["excel_path"])


if __name__ == "__main__":
    main()
